"""Main tables (manuscript/tables.md) and supplementary tables
(manuscript/Supplementary_Information.docx, manuscript/Supplementary_Data.xlsx).
Display-item text 10 pt; supplementary tables left-aligned; italic P.
"""
import re

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font

from common import CACHE, EXPOSURES, NEGATIVE, OUT_LABEL, POSITIVE, PRIMARY, PROJ, RES, SECONDARY, NOVA_DIR, PC

MS = PROJ / "manuscript"
SIG_LAB = {"sig_UPF": "UPF signature", "sig_uPDI": "uPDI signature", "sig_hPDI": "hPDI signature",
           "sig_fibre": "Fibre signature", "sig_BMI": "BMI signature"}
DIET_LAB = {"diet_UPF": "UPF (% energy)", "diet_uPDI": "uPDI", "diet_hPDI": "hPDI", "diet_fibre": "Fibre"}
hr = lambda r: f"{r.hr:.2f} ({r.lci:.2f}–{r.uci:.2f})"


def fp(p):
    return "<0.001" if p < 0.001 else f"{p:.3f}"


# ================================================================== Table 1
person = pd.read_pickle(CACHE / "person.pkl")
sig = pd.read_pickle(CACHE / "signatures.pkl")
o = person.merge(sig, on="eid", how="inner")
S_diet = person[person.n_recalls >= 1]
S_dp = o[(o.consortium == 0) & (o.n_recalls >= 1)]
S_ind = o[(o.consortium == 0) & ~o.train]
cols = {"Self-reported diet cohort": S_diet, "Diet–proteome sample": S_dp, "Independent signature sample": S_ind}


def msd(x, dp=1):
    return f"{x.mean():.{dp}f} ({x.std():.{dp}f})"


def miqr(x, dp=1):
    q = x.quantile([.5, .25, .75])
    return f"{q.iloc[0]:.{dp}f} ({q.iloc[1]:.{dp}f}–{q.iloc[2]:.{dp}f})"


def pct(mask, d):
    return f"{int(mask.sum()):,} ({100 * mask.sum() / len(d):.1f})"


rows = [("Participants, n", {k: f"{len(d):,}" for k, d in cols.items()}),
        ("Age, years", {k: msd(d.age) for k, d in cols.items()}),
        ("Women", {k: pct(d.sex == 0, d) for k, d in cols.items()}),
        ("White ethnicity", {k: pct(d.white == 1, d) for k, d in cols.items()}),
        ("Townsend deprivation index", {k: msd(d.townsend, 2) for k, d in cols.items()}),
        ("College or university degree", {k: pct(d.degree == 1, d) for k, d in cols.items()}),
        ("Current smoker", {k: pct(d.smoking == 2, d) for k, d in cols.items()}),
        ("Alcohol daily or almost daily", {k: pct(d.alcohol_freq == 1, d) for k, d in cols.items()}),
        ("Physical activity, MET-min/week", {k: miqr(d.met_min, 0) for k, d in cols.items()}),
        ("BMI, kg/m²", {k: msd(d.bmi) for k, d in cols.items()}),
        ("Waist circumference, cm", {k: msd(d.waist) for k, d in cols.items()}),
        ("C-reactive protein, mg/L", {k: miqr(d.crp, 2) for k, d in cols.items()}),
        ("Cystatin C, mg/L", {k: msd(d.cysc, 2) for k, d in cols.items()}),
        ("Valid 24-h recalls, n", {k: (miqr(d.n_recalls, 0) if (d.n_recalls >= 1).mean() > .99 else "—")
                                   for k, d in cols.items()}),
        ("UPF, % energy", {k: (msd(d.UPF) if (d.n_recalls >= 1).mean() > .99 else "—") for k, d in cols.items()}),
        ("hPDI, score", {k: (msd(d.hPDI) if (d.n_recalls >= 1).mean() > .99 else "—") for k, d in cols.items()}),
        ("uPDI, score", {k: (msd(d.uPDI) if (d.n_recalls >= 1).mean() > .99 else "—") for k, d in cols.items()}),
        ("Energy-adjusted fibre, g/day", {k: (msd(d.fibre) if (d.n_recalls >= 1).mean() > .99 else "—")
                                          for k, d in cols.items()})]
t1 = pd.DataFrame([{"Characteristic": n, **v} for n, v in rows])
t1.to_csv(RES / "table1.csv", index=False)

# ================================================================== Table 2
cs = pd.read_csv(RES / "cox_signatures.csv")
cd = pd.read_csv(RES / "cox_diet.csv")
t2 = []
for oc in PRIMARY:
    for e in EXPOSURES:
        a = cs[(cs.outcome == oc) & (cs.x == f"sig_{e}")].set_index("model")
        b = cd[(cd.outcome == oc) & (cd.x == f"diet_{e}") & (cd.model == "M1")].iloc[0]
        t2.append({"Outcome": OUT_LABEL[oc] if e == EXPOSURES[0] else "", "Diet dimension": {"UPF": "UPF", "uPDI": "uPDI",
                   "hPDI": "hPDI", "fibre": "Fibre"}[e],
                   "Signature: events/N": f"{int(a.loc['M1', 'events']):,}/{int(a.loc['M1', 'n']):,}" if e == EXPOSURES[0] else "",
                   "Signature, Model 1": hr(a.loc["M1"]), "Signature, Model 3": hr(a.loc["M3"]),
                   "Self-reported: events/N": f"{int(b.events):,}/{int(b.n):,}" if e == EXPOSURES[0] else "",
                   "Self-reported score, Model 1": hr(b)})
t2 = pd.DataFrame(t2)
t2.to_csv(RES / "table2.csv", index=False)


def md_table(d):
    head = "| " + " | ".join(d.columns) + " |\n|" + "---|" * d.shape[1] + "\n"
    return head + "\n".join("| " + " | ".join(str(v) for v in r) + " |" for r in d.itertuples(index=False))


ABBR_MAIN = ("BMI, body mass index; CI, confidence interval; HR, hazard ratio; hPDI, healthful plant-based diet index; "
             "MET, metabolic equivalent of task; UPF, ultra-processed food; uPDI, unhealthful plant-based diet index.")
md = ("# Tables\n\n**Table 1. Characteristics of the three analytic samples**\n\n" + md_table(t1) +
      "\n\nValues are mean (SD), median (IQR) or n (%). The self-reported diet cohort comprises all participants with ≥1 valid "
      "24-h recall; the diet–proteome sample comprises randomly selected, quality-controlled Olink participants with ≥1 valid "
      "recall (used to derive the signatures); the independent signature sample comprises randomly selected Olink participants "
      "outside the signature training set. Diet variables are shown only for samples in which all participants completed a "
      "recall. Abbreviations: " + ABBR_MAIN +
      "\n\n\\pagebreak\n\n**Table 2. Proteomic signatures and self-reported scores of diet and risk of inflammatory bowel "
      "disease, gallstone disease and acute pancreatitis**\n\n" + md_table(t2) +
      "\n\nHRs (95% CIs) per 1-SD higher signature or score. Signature analyses: independent Olink participants (training set "
      "excluded). Self-reported scores: all participants with ≥1 valid 24-h recall. Model 1: age, sex, ethnicity, Townsend index, "
      "education, smoking, alcohol, physical activity, BMI, fasting time, assessment centre, log cystatin C, log creatinine and "
      "log CRP (self-reported scores additionally total energy and number of recalls). Model 3: Model 1 plus waist "
      "circumference, log ALT, log GGT, HbA1c, log triglycerides and the BMI proteomic signature. Non-gallstone acute "
      "pancreatitis excludes participants with gallstones at baseline, censors follow-up at gallstone diagnosis and counts "
      "acute pancreatitis with gallstones diagnosed within 90 days as gallstone-related. Abbreviations: ALT, alanine "
      "aminotransferase; CRP, C-reactive protein; GGT, γ-glutamyltransferase; " + ABBR_MAIN + "\n")
MS.mkdir(parents=True, exist_ok=True)
(MS / "tables.md").write_text(md, encoding="utf-8")

# ================================================================== Supplementary tables
nova = pd.read_csv(NOVA_DIR / "nova_classification.csv", keep_default_na=False)
energy = pd.read_csv(PC / "nova_item_energy.csv").rename(columns={
    "feature": "Feature", "kJ_per_unit": "Estimated kJ per serving (or per g)", "item": "Item", "NOVA": "NOVA group",
    "mean_units_per_recall": "Mean servings per recall", "energy_share_%": "Share of total energy, %"})
perf = pd.read_csv(RES / "signature_performance.csv")
perf_t = pd.DataFrame({"Signature": perf.signature.map({"UPF": "UPF (% energy)", "uPDI": "uPDI", "hPDI": "hPDI",
                                                         "fibre": "Energy-adjusted fibre", "BMI": "BMI"}),
                       "Training, n": perf.n_train.map("{:,}".format), "Testing, n": perf.n_test.map("{:,}".format),
                       "Proteins selected, n": perf.n_proteins, "r in testing set": perf.r_test.round(2)})
top_w = []
for e in EXPOSURES:
    w = pd.read_csv(RES / f"signature_weights_{e}.csv", index_col=0)["weight"]
    w = w.reindex(w.abs().sort_values(ascending=False).index).head(15)
    top_w.append(pd.DataFrame({"Signature": e, "Protein": w.index.str.upper(), "Weight": w.round(3).to_numpy()}))
top_w = pd.concat(top_w)
gut = pd.read_csv(RES / "diet_gutproteins_M1.csv")
gut_t = gut.assign(cell=gut.apply(lambda s: f"{s.beta:+.3f} ({s.lci:+.3f} to {s.uci:+.3f}){'†' if s.bonf < .05 else ''}", axis=1)) \
    .pivot(index="protein", columns="exposure", values="cell")[EXPOSURES].reset_index()
gut_t["protein"] = gut_t.protein.str.upper()
gut_t.columns = ["Protein", "UPF", "uPDI", "hPDI", "Fibre"]


def cox_wide(r, xs, labels, models):
    out = []
    for oc in PRIMARY + ["AP_biliary"] + POSITIVE + SECONDARY + NEGATIVE:
        row = {"Outcome": OUT_LABEL[oc]}
        sub = r[r.outcome == oc]
        if sub.empty:
            continue
        row["Events/N"] = f"{int(sub.events.iloc[0]):,}/{int(sub.n.iloc[0]):,}"
        for x in xs:
            for m in models:
                s = sub[(sub.x == x) & (sub.model == m)]
                if len(s):
                    row[f"{labels[x]}, {m.replace('M', 'Model ')}"] = hr(s.iloc[0])
        out.append(row)
    return pd.DataFrame(out)


s_sig = cox_wide(cs, list(SIG_LAB), SIG_LAB, ["M1", "M2", "M3"])
s_diet = cox_wide(cd, list(DIET_LAB), DIET_LAB, ["M1", "M2"])
sens = pd.read_csv(RES / "cox_sensitivity.csv")
sens["label"] = sens.x.map({**SIG_LAB, **DIET_LAB})
sens_t = []
for (a, m), g in sens.groupby(["analysis", "model"], sort=False):
    for oc in PRIMARY + POSITIVE:
        s = g[g.outcome == oc]
        sens_t.append({"Analysis": f"{a} ({m.replace('M', 'Model ')})", "Outcome": OUT_LABEL[oc],
                       "Events/N": f"{int(s.events.iloc[0]):,}/{int(s.n.iloc[0]):,}",
                       **{s2.label: hr(s2) for s2 in s.itertuples()}})
sens_t = pd.DataFrame(sens_t)
ph = pd.read_csv(RES / "ph_test.csv")
ph_t = ph.assign(x=ph.x.map(SIG_LAB), outcome=ph.outcome.map(OUT_LABEL), ph_p=ph.ph_p.map(fp)) \
    .pivot(index="outcome", columns="x", values="ph_p").reset_index().rename(columns={"outcome": "Outcome"})
med = pd.read_csv(RES / "mediation.csv")
med_t = pd.DataFrame({"Diet": med.exposure.map({"UPF": "UPF", "uPDI": "uPDI", "hPDI": "hPDI", "fibre": "Fibre"}),
                      "Outcome": med.outcome.map(OUT_LABEL), "Mediator": med.mediator.str.upper(),
                      "Total HR (95% CI)": med.apply(lambda s: f"{s.total_HR:.2f} ({s.total_lci:.2f}–{s.total_uci:.2f})", axis=1),
                      "P total": med.total_p.map(fp), "Diet → protein β": med.a_beta.map("{:+.3f}".format),
                      "Protein → outcome HR (95% CI)": med.apply(lambda s: f"{s.b_HR:.2f} ({s.b_lci:.2f}–{s.b_uci:.2f})", axis=1),
                      "Indirect HR (95% CI)": med.apply(lambda s: f"{s.indirect_HR:.3f} ({s.indirect_lci:.3f}–{s.indirect_uci:.3f})", axis=1),
                      "Proportion mediated, %": med.prop_mediated.map(lambda v: "" if pd.isna(v) else f"{100 * v:.1f}")})
fdr = pd.read_csv(RES / "foodgroup_drivers.csv")
dp_all = pd.concat([pd.read_csv(RES / f"diet_proteome_{e}_M1.csv").assign(exposure=e) for e in EXPOSURES])
dp_all = dp_all.rename(columns={"protein": "Protein", "beta": "β", "se": "SE", "p": "P", "lci": "Lower 95% CI",
                                "uci": "Upper 95% CI", "fdr": "FDR-adjusted P", "bonf": "Bonferroni-adjusted P",
                                "exposure": "Diet"})
weights_all = pd.concat([pd.read_csv(RES / f"signature_weights_{e}.csv").assign(Signature=e)
                         for e in EXPOSURES + ["BMI"]]).rename(columns={"Unnamed: 0": "Protein", "weight": "Weight"})

ABBR = ("BMI, body mass index; CI, confidence interval; HR, hazard ratio; hPDI, healthful plant-based diet index; "
        "UPF, ultra-processed food; uPDI, unhealthful plant-based diet index.")
SI = [  # (title, data, note) -> Word
    ("Supplementary Table 1. Performance of the proteomic signatures", perf_t,
     "Elastic net (L1 ratio 0.5; penalty chosen by 5-fold cross-validation) trained in a random 70% of the diet–proteome "
     "sample; r is the Pearson correlation between the signature and the measured exposure in the remaining 30%."),
    ("Supplementary Table 2. Fifteen proteins with the largest absolute weights in each diet signature", top_w,
     "Weights apply to plate-adjusted, standardised NPX values. Full weights are given in Supplementary Data 3."),
    ("Supplementary Table 3. Diet dimensions and intestinal and gut-related plasma proteins", gut_t,
     "β (95% CI) in SD of protein per 1-SD higher diet score, Model 1. †Bonferroni-adjusted P < 0.05 (2,919 proteins)."),
    ("Supplementary Table 4. Proteomic signatures and all outcomes, Models 1–3", s_sig,
     "HR (95% CI) per 1-SD higher signature; independent Olink participants. Model 2: Model 1 plus waist circumference, "
     "log ALT, log GGT, HbA1c and log triglycerides; Model 3: Model 2 plus the BMI proteomic signature."),
    ("Supplementary Table 5. Self-reported diet scores and all outcomes", s_diet,
     "HR (95% CI) per 1-SD higher score in all participants with ≥1 valid 24-h recall."),
    ("Supplementary Table 6. Sensitivity analyses", sens_t,
     "S1 excludes events in the first 2 years; S2 includes UKB-PPP consortium-selected participants; S3 is restricted to "
     "participants who never completed a 24-h recall; S4 restricts self-reported scores to participants with ≥2 valid recalls."),
    ("Supplementary Table 7. Tests of the proportional hazards assumption", ph_t,
     "P values from Schoenfeld-residual tests (rank-transformed time) for the signature term, Model 1."),
    ("Supplementary Table 8. Mediation of diet–outcome associations by intestinal epithelial proteins", med_t,
     "Diet–proteome sample, Model 1. Indirect effect = a × b with 95% CI from 100,000 Monte Carlo draws; proportion mediated "
     "is shown only when the total effect has P < 0.05."),
]

# ---------------------------------------------------------------- tier-1, items 7-8, MR and colocalisation
GS_LAB = {"GS_or_chole": "Gallstone disease or cholecystectomy", "GS_hosp": "Hospital-diagnosed gallstone disease",
          "Cholecystectomy": "Cholecystectomy"}
LAB = {**OUT_LABEL, **GS_LAB}
t1x = pd.read_csv(RES / "tier1_extended_adjustment.csv")
s9 = []
for oc in ["Gallstones", "GS_or_chole", "GS_hosp", "Cholecystectomy", "IBD", "Crohn", "UC", "Acute_pancreatitis", "AP_nonbiliary"]:
    for m, ml in [("M1", "Model 1"), ("M1x", "Model 1 + gallstone-related factors"), ("M3x", "Model 3 + gallstone-related factors")]:
        g = t1x[(t1x.outcome == oc) & (t1x.model == m)].set_index("x")
        s9.append({"Outcome": LAB[oc], "Model": ml, "Events/N": f"{int(g.events.iloc[0]):,}/{int(g.n.iloc[0]):,}",
                   **{SIG_LAB[x]: hr(g.loc[x]) for x in g.index}})
s9 = pd.DataFrame(s9)
dr = pd.read_csv(RES / "tier1_dose_response.csv")
s10 = dr.assign(**{"Outcome": dr.outcome.map(LAB), "Signature": dr.x.map(SIG_LAB), "Quintile": dr.quintile.map(lambda q: f"Q{q}"),
                   "Events/person-years": dr.apply(lambda r: f"{int(r.events)}/{r.person_years:,.0f}", axis=1),
                   "Rate per 1,000 person-years": dr.rate_per_1000py.round(2), "10-year cumulative incidence, %": dr["cuminc_10y_%"].round(2),
                   "HR (95% CI)": dr.apply(lambda r: "1.00 (reference)" if r.quintile == 1 else f"{r.hr:.2f} ({r.lci:.2f}–{r.uci:.2f})", axis=1),
                   "P for trend": dr.apply(lambda r: fp(r.p_trend) if r.quintile == 1 else "", axis=1)})[
    ["Outcome", "Signature", "Quintile", "Events/person-years", "Rate per 1,000 person-years", "10-year cumulative incidence, %",
     "HR (95% CI)", "P for trend"]]
md_ = pd.read_csv(RES / "tier1_missing_data.csv")
ev = t1x[t1x.model == "M1x"].set_index(["outcome", "x"])
s11 = []
for oc in ["IBD", "Crohn", "UC", "Gallstones", "Acute_pancreatitis", "AP_nonbiliary"]:
    for x in SIG_LAB:
        if x == "sig_BMI":
            continue
        cc = md_[(md_.analysis == "complete case") & (md_.outcome == oc) & (md_.x == x)].iloc[0]
        mi = md_[(md_.analysis == "multiple imputation") & (md_.outcome == oc) & (md_.x == x)].iloc[0]
        e = ev.loc[(oc, x)]
        s11.append({"Outcome": LAB[oc], "Signature": SIG_LAB[x], "Complete case, HR (95% CI)": hr(cc),
                    "Multiple imputation, HR (95% CI)": hr(mi), "E-value (point)": f"{e.evalue_point:.2f}",
                    "E-value (CI limit)": f"{e.evalue_ci:.2f}"})
s11 = pd.DataFrame(s11)
rp = pd.read_csv(RES / "residual_signature_performance.csv")
rc7 = pd.read_csv(RES / "residual_signature_cox.csv")
RS = {"rsig_UPF": "UPF", "rsig_uPDI": "uPDI", "rsig_hPDI": "hPDI", "rsig_fibre": "Fibre"}
s12 = []
for oc in ["IBD", "Crohn", "UC", "Gallstones", "Acute_pancreatitis", "AP_nonbiliary", "MASLD", "Cirrhosis"] + NEGATIVE:
    for m in ("M1", "M3"):
        g = rc7[(rc7.outcome == oc) & (rc7.model == m)].set_index("x")
        s12.append({"Outcome": LAB[oc], "Model": m.replace("M", "Model "), **{RS[x]: hr(g.loc[x]) for x in RS}})
s12 = pd.DataFrame(s12)
s12_perf = ("Residual signatures retained " + ", ".join(f"{RS['rsig_' + r.signature]} {int(r.n_proteins)} proteins (r = {r.r_test_residual_diet:.2f} "
            f"with the residualised score in the test set; r = {r.r_with_original_signature:.2f} with the original signature; "
            f"r = {r.r_with_BMI_signature:.2f} with the BMI signature)" for r in rp.itertuples()) + ". Covariates explained "
            f"{100 * rp.r2_diet_explained_by_covariates.min():.0f}%–{100 * rp.r2_diet_explained_by_covariates.max():.0f}% of the variance of the diet scores.")
st = pd.read_csv(RES / "repeat_stability.csv")
s13 = pd.DataFrame({"Measure": st.signature.map(lambda v: {"UPF": "UPF reduced signature", "uPDI": "uPDI reduced signature",
                    "hPDI": "hPDI reduced signature", "fibre": "Fibre reduced signature"}.get(v, v)),
                    "Proteins, n": st.n_proteins.map(lambda v: "" if pd.isna(v) else int(v)),
                    "r with full signature": st.r_with_full_signature.map(lambda v: "" if pd.isna(v) else f"{v:.2f}"),
                    "Spearman r, baseline vs imaging visit 1": st.spearman_baseline_vs_instance2.map("{:.2f}".format),
                    "Spearman r, baseline vs imaging visit 2": st.spearman_baseline_vs_instance3.map("{:.2f}".format)})
s13_note = (f"Repeat Olink measurements (1,463 proteins) were available for {int(st.n_instance2.max())} and {int(st.n_instance3.max())} "
            f"participants of the baseline proteomics sample at the first and second imaging visits (median "
            f"{st.years_to_instance2.max():.1f} and {st.years_to_instance3.max():.1f} years after baseline). Reduced signatures were "
            "retrained on the proteins measured at both visits.")
mrf = pd.read_csv(RES / "mr_guca2a_full.csv")
mro = pd.read_csv(RES / "mr_results.csv")
s14 = pd.concat([
    pd.DataFrame({"Protein": "GUCA2A", "Outcome (FinnGen R12)": mrf.outcome, "Instrument(s)": mrf.instrument,
                  "Variant(s)": mrf.rsid.astype(str).str.replace(",", "/"), "SNPs, n": mrf.nsnp,
                  "OR per SD (95% CI)": mrf.apply(lambda r: f"{r.OR:.2f} ({r.lci:.2f}–{r.uci:.2f})", axis=1), "P": mrf.p.map(fp)}),
    pd.DataFrame({"Protein": mro.protein, "Outcome (FinnGen R12)": mro.outcome, "Instrument(s)": "cis-pQTL (UKB-PPP Supplementary Table 16)",
                  "Variant(s)": "", "SNPs, n": mro.nsnp, "OR per SD (95% CI)": mro.apply(lambda r: f"{r.OR:.2f} ({r.lci:.2f}–{r.uci:.2f})", axis=1),
                  "P": mro.p.map(fp)})[lambda t: t.Protein != "GUCA2A"]])
co = pd.read_csv(RES / "coloc_guca2a.csv")
s15 = pd.DataFrame({"Outcome (FinnGen R12)": co.outcome, "Prior p12": co.p12.map(lambda v: f"{v:.0e}"), "SNPs, n": co.nsnps,
                    **{f"PP.H{i}": co[f"PP.H{i}"].map(lambda v: f"{v:.3f}") for i in range(5)},
                    "Lowest outcome P in region": co.outcome_lead_p_in_region.map(lambda v: f"{v:.1e}")})
SI += [
    ("Supplementary Table 9. Further adjustment for gallstone-related factors and alternative gallstone definitions", s9,
     "HR (95% CI) per 1-SD higher signature; independent Olink participants. Gallstone-related factors: hormone replacement "
     "therapy, oral contraceptive use, number of live births (women), cholesterol-lowering, antihypertensive and insulin treatment, "
     "doctor-diagnosed diabetes and weight change in the previous year. Cholecystectomy: OPCS-4 J18."),
    ("Supplementary Table 10. Dose–response: quintiles of the proteomic signatures", s10,
     "Model 1 plus gallstone-related factors. 10-year cumulative incidence by Kaplan–Meier; P for trend from the quintile median "
     "modelled as a continuous variable."),
    ("Supplementary Table 11. Complete-case and multiple-imputation analyses and E-values", s11,
     "Model 1 plus gallstone-related factors. Multiple imputation by chained equations (5 imputations, Rubin's rules). E-values "
     "refer to the main analysis; an E-value of 1.00 for the CI limit indicates a CI that includes 1."),
    ("Supplementary Table 12. Diet-specific (residual) proteomic signatures and outcomes", s12,
     "Each diet score was residualised on age, sex, ethnicity, Townsend index, education, smoking, alcohol, physical activity, BMI, "
     "waist circumference, ALT, GGT, HbA1c, triglycerides, cystatin C, creatinine, CRP, energy intake, number of recalls, fasting "
     "time and centre before training the elastic net. " + s12_perf),
    ("Supplementary Table 13. Within-person stability of the signatures and intestinal proteins", s13, s13_note),
    ("Supplementary Table 14. Two-sample cis-pQTL Mendelian randomization", s14,
     "Exposure: plasma protein (per SD) from UKB-PPP (European discovery cohort); outcomes from FinnGen R12. For GUCA2A, marginal "
     "per-allele estimates from the full summary statistics; IVW over SuSiE-independent cis signals. For other proteins, "
     "Wald ratios with the cis-pQTL matched in FinnGen (per-allele estimates derived from Supplementary Table 16 of the UKB-PPP paper)."),
    ("Supplementary Table 15. Colocalisation of the GUCA2A cis-pQTL with digestive outcomes", s15,
     "coloc.abf within ±1 Mb of GUCA2A; H3, distinct causal variants; H4, a shared causal variant."),
]


# ---------------------------------------------------------------- tier 3 and NOVA sensitivity
sg = pd.read_csv(RES / "tier3_subgroups.csv")
s16 = []
for (oc, gname), g in sg.groupby(["outcome", "subgroup"], sort=False):
    for lvl, gg in g.groupby("level", sort=False):
        gg = gg.set_index("x")
        s16.append({"Outcome": LAB[oc], "Subgroup": f"{gname}: {lvl}", "Events/N": f"{int(gg.events.iloc[0]):,}/{int(gg.n.iloc[0]):,}",
                    **{SIG_LAB[x]: hr(gg.loc[x]) for x in gg.index}})
    s16.append({"Outcome": LAB[oc], "Subgroup": f"{gname}: P for interaction", "Events/N": "",
                **{SIG_LAB[x]: fp(v) for x, v in g.groupby("x").p_interaction.first().items()}})
s16 = pd.DataFrame(s16)
pr = pd.read_csv(RES / "tier3_prediction.csv")
s17 = pd.DataFrame({"Outcome": pr.outcome.map(LAB), "Model": pr.model.map(lambda m: m.replace("+ sig_", "+ ").replace("fibre", "Fibre") + (" signature" if m.startswith("+ sig_") else "")),
                    "C-index": pr.C.map("{:.3f}".format),
                    "Change in C-index (95% CI)": pr.apply(lambda r: "Reference" if pd.isna(r.delta_C) else f"{r.delta_C:+.3f} ({r.lci:+.3f} to {r.uci:+.3f})", axis=1)})
sb = pd.read_csv(RES / "tier3_substitution.csv")
s18 = pd.DataFrame({"Outcome": sb.outcome.map(LAB), "Analysis": sb.analysis, "Events/N": sb.apply(lambda r: f"{int(r.events):,}/{int(r.n):,}", axis=1),
                    "HR (95% CI)": sb.apply(hr, axis=1), "P": sb.p.map(fp)})
fg = pd.read_csv(RES / "tier3_finegray.csv")
s19 = []
for oc, g in fg.groupby("outcome", sort=False):
    g = g.set_index("x")
    s19.append({"Outcome": LAB[oc], "Events/deaths/N": f"{int(g.events.iloc[0]):,}/{int(g.deaths.iloc[0]):,}/{int(g.n.iloc[0]):,}",
                **{SIG_LAB[x]: f"{r.shr:.2f} ({r.lci:.2f}–{r.uci:.2f})" for x, r in g.iterrows()}})
s19 = pd.DataFrame(s19)
ns = pd.read_csv(RES / "nova_sensitivity.csv")
nsi = pd.read_csv(RES / "nova_sensitivity_info.csv").set_index("scenario")
SCN = {"main": "Main classification", "upper": "All uncertain items as UPF", "lower": "No uncertain item as UPF"}
s20 = []
for (an, oc), g in ns.groupby(["analysis", "outcome"], sort=False):
    g = g.set_index("scenario")
    s20.append({"Analysis": an.replace("M1", "Model 1").replace("M3", "Model 3"), "Outcome": LAB[oc],
                **{SCN[k]: hr(g.loc[k]) for k in SCN if k in g.index}})
s20 = pd.DataFrame(s20)
s20_note = ("Mean UPF intake (% of energy): " + "; ".join(f"{SCN[k].lower()} {nsi.loc[k, 'mean_UPF_%E']:.1f}" for k in SCN) +
            ". Correlation with the main UPF estimate: " + "; ".join(f"{SCN[k].lower()} r = {nsi.loc[k, 'r_with_main_UPF']:.2f}" for k in ("upper", "lower")) +
            ". Test-set r of the retrained UPF signature: " + "; ".join(f"{SCN[k].lower()} {nsi.loc[k, 'r_test']:.2f}" for k in SCN) + ".")
SI += [
    ("Supplementary Table 16. Subgroup analyses", s16,
     "Model 1 plus gallstone-related factors within strata; P for interaction from Wald tests of signature-by-subgroup product terms."),
    ("Supplementary Table 17. Added predictive value of the signatures for gallstone disease and IBD", s17,
     "Harrell's C from 5-fold cross-validated Cox models; conventional risk factors: age, sex, ethnicity, Townsend index, BMI, waist "
     "circumference, smoking, alcohol, physical activity, triglycerides, HbA1c and gallstone-related factors. 95% CI from 500 bootstrap samples."),
    ("Supplementary Table 18. Isocaloric substitution of UPF with unprocessed or minimally processed foods", s18,
     "Self-reported diet cohort, Model 1 plus gallstone-related factors. Substitution model: energy from NOVA 1 foods and from all "
     "other non-UPF foods (per 10% of energy) entered together with total energy, UPF omitted."),
    ("Supplementary Table 19. Fine–Gray competing-risk models with death as the competing event", s19,
     "Subdistribution HRs (95% CI) per 1-SD higher signature, Model 1 plus gallstone-related factors."),
    ("Supplementary Table 20. Sensitivity of UPF results to the NOVA classification of uncertain items", s20, s20_note),
]

DATA = [("Supplementary Data 1. Item-level NOVA classification of Oxford WebQ items", nova),
        ("Supplementary Data 2. Estimated energy per serving from non-negative least-squares regression", energy.round(3)),
        ("Supplementary Data 3. Proteomic signature weights", weights_all),
        ("Supplementary Data 4. Diet dimensions and 2,919 plasma proteins (Model 1)", dp_all),
        ("Supplementary Data 5. Food groups and selected intestinal proteins (Model 1)", fdr)]

# ---------------------------------------------------------------- Word writer
P_RE = re.compile(r"(?<![A-Za-z])P(?![A-Za-z])")
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)


def add_text(par, text, bold=False, italic=False):
    pos = 0
    for m in list(P_RE.finditer(text)) + [None]:
        end = m.start() if m else len(text)
        if end > pos:
            r = par.add_run(text[pos:end]); r.font.size = Pt(10); r.bold = bold or None; r.italic = italic or None
        if m:
            r = par.add_run("P"); r.font.size = Pt(10); r.bold = bold or None; r.italic = True; pos = m.end()


def border(cell, edges):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for e in ("top", "bottom", "left", "right"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single" if e in edges else "nil")
        if e in edges:
            el.set(qn("w:sz"), "8"); el.set(qn("w:color"), "000000")
        b.append(el)
    tcPr.append(b)


first = True
for title, d, note in SI:
    land = d.shape[1] > 5 or d.astype(str).apply(lambda c: c.str.len().max()).sum() > 95
    sec = doc.sections[0] if first else doc.add_section()
    first = False
    if land and sec.page_width < sec.page_height:
        sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
    if not land and sec.page_width > sec.page_height:
        sec.orientation = WD_ORIENT.PORTRAIT; sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for mg in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, mg, Cm(1.6))
    add_text(doc.add_paragraph(), title, bold=True)
    t = doc.add_table(rows=len(d) + 1, cols=d.shape[1])
    for j, h in enumerate(d.columns):
        c = t.cell(0, j); c.text = ""; add_text(c.paragraphs[0], str(h), bold=True); border(c, ("top", "bottom"))
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    for i, row in enumerate(d.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""
            add_text(c.paragraphs[0], "" if (not isinstance(v, str) and pd.isna(v)) else str(v))
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            border(c, ("bottom",) if i == len(d) else ())
    add_text(doc.add_paragraph(), note + " Abbreviations: " + ABBR, italic=False)
doc.save(MS / "Supplementary_Information.docx")

# ---------------------------------------------------------------- Excel writer
F10, F10B, LEFT = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True), Alignment(horizontal="left")


def xl(v, bold=False):
    if not isinstance(v, str) or not P_RE.search(v):
        return v
    f = lambda it: InlineFont(rFont="Arial", sz=10, b=bold or None, i=it or None)
    out, pos = [], 0
    for m in P_RE.finditer(v):
        if m.start() > pos:
            out.append(TextBlock(f(False), v[pos:m.start()]))
        out.append(TextBlock(f(True), "P")); pos = m.end()
    if pos < len(v):
        out.append(TextBlock(f(False), v[pos:]))
    return CellRichText(out)


wb = Workbook(); ws0 = wb.active; ws0.title = "Contents"
for i, (title, _) in enumerate(DATA, start=1):
    ws0.cell(row=i, column=1, value=title).font = F10
ws0.column_dimensions["A"].width = 100
for title, d in DATA:
    ws = wb.create_sheet(title.split(".")[0].replace("Supplementary ", "Suppl. "))
    ws.cell(row=1, column=1, value=title).font = F10B
    for j, c in enumerate(d.columns, start=1):
        h = ws.cell(row=3, column=j, value=xl(str(c), bold=True)); h.font, h.alignment = F10B, LEFT
    for i, row in enumerate(d.itertuples(index=False), start=4):
        for j, v in enumerate(row, start=1):
            v = None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v)
            c = ws.cell(row=i, column=j, value=v); c.font, c.alignment = F10, LEFT
    for j, c in enumerate(d.columns, start=1):
        w = max([len(str(c))] + [len(str(v)) for v in d.iloc[:200, j - 1]]) + 2
        ws.column_dimensions[ws.cell(row=3, column=j).column_letter].width = min(60, max(9, w))
wb.save(MS / "Supplementary_Data.xlsx")
print(t1.to_string(index=False))
print(t2.to_string(index=False))
